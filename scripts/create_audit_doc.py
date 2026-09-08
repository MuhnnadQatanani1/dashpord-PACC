from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "docs/PACC-Portal-Audit.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tbl_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_status_table(doc):
    rows = [
        ("Homepage map and API-key warning", "Addressed locally", "No-label basemap is configured; production deployment still needs verification."),
        ("Gaza and geographic analysis", "Confirmed direction", "Gaza remains on the map and in the governorate complaint data by decision."),
        ("Contact and footer links", "Addressed", "PACC contact and website links are present."),
        ("Control panel", "Partially addressed", "Hidden route, server sessions, protected mutations, report upload, and Analytics settings are implemented."),
        ("PDF and print output", "Addressed", "The feedback component is excluded from print output."),
        ("Law-enforcement units", "Addressed", "Count and percentage indicators use the requested units."),
        ("Accessibility", "Partially addressed", "Core controls exist; full automated and screen-reader audit remains."),
        ("Arabic TTS", "Locally verified", "Python TTS generated Arabic WAV audio with the installed Majed voice."),
        ("English version", "Pending committee review", "Translation keys are synchronized; language quality approval remains with the committee."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for cell, text in zip(table.rows[0].cells, ("Area", "Status", "Evidence or note")):
        shade(cell, "1F4E79")
        set_cell_text(cell, text, bold=True, color="FFFFFF")
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            if index % 2 == 1:
                shade(cell, "F2F6FA")
            set_cell_text(cell, text)
    doc.add_paragraph()


def add_finding(doc, title, body, severity="High"):
    p = doc.add_paragraph(style="Heading 3")
    p.add_run(f"{severity}  {title}")
    doc.add_paragraph(body)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph(style="Title")
    title.add_run("PACC Portal Audit")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("Current implementation review and release readiness assessment")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph("Audit scope: the portal repository, local implementation, control panel foundation, geographic analysis, accessibility features, text-to-speech, and Arabic and English interface coverage.")
    doc.add_paragraph("Overall conclusion: the portal has a solid working foundation and the requested map direction is now clear, but it is not ready for final production approval until the remaining security hardening and deployment verification are completed.")

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("The portal currently supports Arabic and English content, geographic complaint visualization, report browsing, accessibility controls, Arabic text-to-speech, and a hidden report-management route. Gaza remains intentionally included in the map and geographic complaint data. The first control-panel iteration now includes server-side admin sessions, protected report mutations, and Google Analytics settings.")
    doc.add_paragraph("The most important remaining concerns are password hashing, CSRF protection for administrator actions, access control for unpublished report files, and verification that Python Arabic TTS works in the deployed runtime. Full linting also remains outstanding because the repository contains formatting and code-quality errors.")

    doc.add_heading("Implementation Status", level=1)
    add_status_table(doc)

    doc.add_heading("Priority Findings", level=1)
    add_finding(doc, "Administrator passwords use unsalted SHA-256", "The database layer hashes administrator passwords with plain SHA-256. This is not appropriate for password storage because it is fast and unsalted. Replace it with Argon2id or bcrypt before production use. Reference: src/lib/db.server.ts lines 11–12.", "High")
    add_finding(doc, "CSRF protection is still required", "Administrator report and analytics operations use cookie-authenticated POST requests. SameSite cookies help, but the write operations should also validate a CSRF token or a trusted Origin header. References: src/lib/reports.functions.ts lines 436, 464, and 507; src/lib/analytics.functions.ts line 39.", "High")
    add_finding(doc, "Unpublished report files require an access check", "The file retrieval function selects a report file by ID without checking publication status or administrator authentication. An unpublished file should not be retrievable through a public endpoint. Reference: src/lib/reports.functions.ts lines 541–550.", "High")
    add_finding(doc, "Production TTS runtime is not confirmed", "The server spawns Python and relies on pyttsx3 plus an operating-system voice. Arabic audio works locally with the Majed voice, but the deployed runtime may not provide Python, pyttsx3, or an Arabic system voice. Production verification or a hosted TTS provider is required. References: src/server.ts lines 80–113 and scripts/tts.py.", "High")
    add_finding(doc, "Google Analytics requires privacy and production verification", "The control panel validates and stores the Measurement ID and loads tracking only when enabled. Before production activation, confirm the organization’s privacy and consent requirements and verify that the deployed site sends data to the intended property.", "Medium")
    add_finding(doc, "Full linting is not clean", "The production build passes, and targeted linting for the new control-panel and analytics files passes. The full repository lint currently reports 44 errors and 7 warnings, mainly formatting and existing code-quality issues. These should be resolved before release approval.", "Medium")

    doc.add_heading("Control Panel Review", level=1)
    doc.add_paragraph("The current panel is available at /reports/manage and is intentionally omitted from public navigation. It supports administrator login, report upload, report editing, report deletion, and Google Analytics settings. Server-side session cookies are used for protected operations, and the previous hard-coded default administrator password was removed.")
    doc.add_paragraph("The panel should receive a final security pass before it is shared with certified users. Recommended controls include Argon2id or bcrypt password storage, CSRF protection, login rate limiting, administrator password reset procedures, session cleanup, role separation, and audit logging for report and analytics changes.")

    doc.add_heading("Accessibility and TTS Review", level=1)
    doc.add_paragraph("The portal includes a skip link, keyboard focus styling, an accessibility settings widget, contrast controls, font-size controls, spacing controls, link highlighting, image hiding, motion reduction, and localized ARIA labels. These are good foundations, but a final audit with keyboard-only navigation, a screen reader, and automated checks is still required.")
    doc.add_paragraph("The Arabic text-to-speech pipeline successfully generated a valid WAV file locally using pyttsx3 and the installed Arabic Majed voice. The implementation normalizes Arabic text, expands common abbreviations, splits long content, and reads content in segments. The remaining risk is deployment portability rather than local functionality.")

    doc.add_heading("Language and Content Review", level=1)
    doc.add_paragraph("Arabic and English translation keys are synchronized. The English navigation label for the complaint link is now Make a Complaint, while the Arabic label remains تقديم بلاغ. The English content is ready for Translation Committee review, but that committee review is still the approval step for wording, terminology, and tone.")
    doc.add_paragraph("The geographic analysis intentionally keeps Gaza visible, including governorate complaint counts. The map background uses a no-label style to avoid unwanted Hebrew basemap labels, while the portal’s own Arabic and English labels remain available.")

    doc.add_heading("Verification Evidence", level=1)
    evidence = [
        ("Production build", "Pass"),
        ("Targeted lint for new control-panel and analytics files", "Pass"),
        ("Full repository lint", "Fail: 44 errors and 7 warnings"),
        ("Local Arabic TTS WAV generation", "Pass"),
        ("Local map visual verification", "Pass"),
        ("Live Vercel verification of the new panel", "Pending deployment"),
        ("Full screen-reader and automated accessibility audit", "Pending"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    for cell, text in zip(table.rows[0].cells, ("Check", "Result")):
        shade(cell, "1F4E79")
        set_cell_text(cell, text, bold=True, color="FFFFFF")
    for index, row in enumerate(evidence):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            if index % 2 == 1:
                shade(cell, "F2F6FA")
            set_cell_text(cell, text)

    doc.add_heading("Release Recommendations", level=1)
    for item in [
        "Replace SHA-256 password hashing with Argon2id or bcrypt.",
        "Add CSRF protection and login rate limiting to administrator operations.",
        "Restrict unpublished report file retrieval.",
        "Verify the control panel against the production database and configured administrator credentials.",
        "Verify Google Analytics activation, consent requirements, and the correct Measurement ID in production.",
        "Test Arabic TTS on the deployed runtime or move to a hosted TTS provider.",
        "Resolve full-lint errors and complete the accessibility review.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("PACC Portal Audit")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(128, 128, 128)

    doc.save(OUT)


if __name__ == "__main__":
    main()
